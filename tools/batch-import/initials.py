# Initials in the 2026 batch reports (docs/batch-report/import-2026.md).
# Usage: python3 -I initials.py <out.xlsx> <users.json> <workbook.xlsx>...
# users.json: [{"name": ..., "initials": ...}] from the live users table.
# Cells are found by their labels, not by fixed rows, as layouts shift.
import sys, re, json, os, openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter, column_index_from_string

TOKEN = re.compile(r'^[A-Za-z]{1,4}[0-9]?$')
SPLIT = re.compile(r'[/,&+;]|\s+|\ben\b', re.I)


def text(v):
    return v.strip() if isinstance(v, str) else ''


def find(ws, col, pattern):
    """Rows where column `col` holds a label matching `pattern`."""
    rx = re.compile(pattern, re.I)
    return [c.row for c in ws[col] if rx.search(text(c.value))]


def cells(ws, rows, cols):
    for r in rows:
        for c in cols:
            yield ws[f'{c}{r}']


def col_range(a, b):
    return [get_column_letter(i) for i in range(column_index_from_string(a), column_index_from_string(b) + 1)]


def person_cells(wb):
    """(sheet, cell, role) for every cell that names a person."""
    out = []
    if 'Brouwen' in wb.sheetnames:
        ws = wb['Brouwen']
        # One column pair per brew from S: the value column holds the brewer,
        # the unit column sometimes a second brewer.
        for r in find(ws, 'P', r'^naam brouwer'):
            for i, c in enumerate(col_range('S', 'AJ')):
                out.append((ws, f'{c}{r}', f'brewer (brew {i // 2 + 1})'))
        for r in find(ws, 'K', r'^door:?$'):
            out.append((ws, f'N{r}', 'yeast pitched by'))
    if 'Gistkaart Invoer' in wb.sheetnames:
        ws = wb['Gistkaart Invoer']
        for c in ws[12]:
            if text(c.value).lower() == 'operator':
                for r in range(13, ws.max_row + 1):
                    if r >= 38 and c.column > 23:  # plan / calculator block
                        break
                    out.append((ws, f'{c.column_letter}{r}', 'fermentation log operator'))
    if 'Verwerking' in wb.sheetnames:
        ws = wb['Verwerking']
        for r in find(ws, 'A', r'^name operator start'):
            out.append((ws, f'D{r}', 'processing operator start'))
        for r in find(ws, 'A', r'^name operator end'):
            out.append((ws, f'D{r}', 'processing operator end'))
    if 'Afvullen' in wb.sheetnames:
        ws = wb['Afvullen']
        labels = {r: text(ws[f'I{r}'].value).lower() for r in range(1, 24) if text(ws[f'I{r}'].value)}
        roles = [('afgevuld door', 'filled by'), ('monstername door', 'sampled by'),
                 ('controle 1e verpakking', 'first package check')]
        for r, lab in labels.items():
            for key, role in roles:
                if lab.startswith(key):
                    # The label row plus following rows without a label of their own.
                    rr = r
                    while rr == r or (rr not in labels and rr < r + 3):
                        for c in cells(ws, [rr], col_range('M', 'R')):
                            out.append((ws, c.coordinate, role))
                        rr += 1
    if 'Afvulverslag' in wb.sheetnames:
        ws = wb['Afvulverslag']
        for r in find(ws, 'B', r'^ingevuld door'):
            for c in col_range('D', 'I'):
                out.append((ws, f'{c}{r}', 'filling report completed by'))
        for r in find(ws, 'B', r'^ingevoerd in mdrank'):
            for c in col_range('D', 'I'):
                out.append((ws, f'{c}{r}', 'entered in Mdrank by'))
    return out


def main():
    out_path, users_path, books = sys.argv[1], sys.argv[2], sys.argv[3:]
    users = json.load(open(users_path))
    by_initials = {}
    for u in users:
        if u.get('initials'):
            by_initials.setdefault(u['initials'].upper(), []).append(u['name'])

    occ = []  # batch, sheet, cell, role, raw, initial, note
    for p in books:
        wb = openpyxl.load_workbook(p, data_only=True)
        batch = os.path.basename(p)[:5]
        for ws, coord, role in person_cells(wb):
            raw = ws[coord].value
            if raw is None or not str(raw).strip():
                continue
            raw = str(raw).strip()
            parts = [t for t in SPLIT.split(raw) if t and t.strip()]
            if parts and all(TOKEN.match(t) for t in parts):
                for t in parts:
                    note = 'lower case in workbook' if t != t.upper() else ''
                    occ.append((batch, ws.title, coord, role, raw, t.upper(), note))
            else:
                occ.append((batch, ws.title, coord, role, raw, raw, 'not initials: check'))

    summary = {}
    for batch, sheet, coord, role, raw, ini, note in occ:
        s = summary.setdefault(ini, {'n': 0, 'batches': set(), 'roles': set()})
        s['n'] += 1
        s['batches'].add(batch)
        s['roles'].add(role.split(' (')[0])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Initials'
    head = ['Initials', 'Name (fill in)', 'Email (fill in)', 'User in app with these initials',
            'Times used', 'Batches', 'Used as']
    ws.append(head)
    for ini in sorted(summary):
        s = summary[ini]
        ws.append([ini, '', '', ', '.join(by_initials.get(ini, [])), s['n'],
                   ', '.join(sorted(s['batches'])), ', '.join(sorted(s['roles']))])
    fill = PatternFill('solid', fgColor='FFF2CC')
    for r in range(2, ws.max_row + 1):
        ws[f'B{r}'].fill = fill
        ws[f'C{r}'].fill = fill
    ws2 = wb.create_sheet('Where used')
    ws2.append(['Batch', 'Sheet', 'Cell', 'Field', 'Value in workbook', 'Initials', 'Note'])
    for row in occ:
        ws2.append(list(row))
    for sh in (ws, ws2):
        for c in sh[1]:
            c.font = Font(bold=True)
        sh.freeze_panes = 'A2'
        for col in sh.columns:
            w = max(len(str(c.value or '')) for c in col)
            sh.column_dimensions[col[0].column_letter].width = min(max(10, w + 2), 60)
    wb.save(out_path)
    print(f'{len(books)} workbooks, {len(occ)} entries, {len(summary)} unique')
    for ini in sorted(summary):
        print(f"  {ini:6} {summary[ini]['n']:3}  {', '.join(by_initials.get(ini, [])) or '-':28} {', '.join(sorted(summary[ini]['roles']))}")


if __name__ == '__main__':
    main()
