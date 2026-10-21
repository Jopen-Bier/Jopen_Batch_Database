# Reader for the 2026 Excel batch reports (docs/batch-report/import-2026.md).
# Usage: python3 -I reader.py <out_dir> <ref.json> <workbook.xlsx>...
# Writes <out_dir>/<batch>.json per workbook: the rows for the app's tables in
# its column names, people as initials, tanks / tags / articles / materials as
# codes (resolved to ids by the SQL generator), plus a review list (source
# cell, label, value, target, status) and warnings.
# ref.json: {"users": [{"initials"}], "tanks": [{"code", "type", "n"}],
#            "tags": [{"code"}], "articles": [{"code", "fmt"}], "materials": [{"code"}]}
# Cells are found by their labels where layouts differ between workbooks.
# Formula and error cells are not imported (the app calculates its own values),
# except the yeast block, whose formulas carry the recipe values the brewer
# accepted (status "formula").
import sys, os, re, json, datetime
from zoneinfo import ZoneInfo
import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string

TZ = ZoneInfo('Europe/Amsterdam')
ERRORS = ('#VALUE!', '#N/A', '#DIV/0!', '#REF!', '#NAME?', '#NUM!', '#NULL!')
PEOPLE_SPLIT = re.compile(r'[/,&+;]|\s+|\ben\b', re.I)


def col(c, offset=0):
    return get_column_letter(column_index_from_string(c) + offset)


def norm(s):
    return re.sub(r'\s+', ' ', str(s or '')).strip().lower().rstrip(':').strip()


MAX_COL = 40  # nothing is read right of column AN (the old reports are 16000 columns wide)


class Cell:
    __slots__ = ('value',)

    def __init__(self, value):
        self.value = value


class Sheet:
    """The cells of one sheet, read once (streaming), addressed like openpyxl: ws['B7'].value."""

    def __init__(self, ws):
        self.cells = {}
        self.max_row = 0
        for row in ws.iter_rows(max_col=MAX_COL):
            for c in row:
                if c.value is not None and getattr(c, 'coordinate', None):
                    self.cells[c.coordinate] = c.value
                    self.max_row = max(self.max_row, c.row)

    def __getitem__(self, ref):
        return Cell(self.cells.get(ref))


def read_sheets(path, data_only):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=data_only)
    try:
        return {ws.title: Sheet(ws) for ws in wb.worksheets}
    finally:
        wb.close()


class Book:
    """Values (as last calculated) and formulas of one workbook, plus the review."""

    def __init__(self, path):
        self.v = read_sheets(path, True)
        self.f = read_sheets(path, False)
        self.v_names = list(self.v)
        self.review = []
        self.warnings = []

    def has(self, sheet):
        return sheet in self.v

    def raw(self, sheet, ref):
        return self.v[sheet][ref].value

    def is_formula(self, sheet, ref):
        fv = self.f[sheet][ref].value
        return isinstance(fv, str) and fv.startswith('=')

    def get(self, sheet, ref, allow_formula=False):
        """(value, status): None for empty, formula (unless allowed) or error cells."""
        v = self.raw(sheet, ref)
        if v is None or (isinstance(v, str) and not v.strip()):
            return None, 'empty'
        if isinstance(v, str) and v.strip() in ERRORS:
            return None, 'error'
        if self.is_formula(sheet, ref):
            return (v, 'formula') if allow_formula else (None, 'formula')
        return (v.strip() if isinstance(v, str) else v), 'read'

    def note(self, sheet, ref, label, value, target, status, note=''):
        if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
            value = value.isoformat()
        self.review.append({'sheet': sheet, 'cell': ref, 'label': label, 'value': value,
                            'target': target, 'status': status, 'note': note})

    def warn(self, sheet, ref, text):
        self.warnings.append(f'{sheet}!{ref}: {text}')

    def rows(self, sheet, column):
        """(row, normalized label) for every text label in a column."""
        ws = self.v[sheet]
        out = []
        for r in range(1, ws.max_row + 1):
            v = ws[f'{column}{r}'].value
            if isinstance(v, str) and v.strip():
                out.append((r, norm(v)))
        return out


# ---------------------------------------------------------------- converters

def to_num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(',', '.')
    if s in ('-', ''):
        return None
    try:
        return float(s)
    except ValueError:
        raise ValueError(v)


def to_bool(v):
    s = norm(v)
    if s in ('ja', 'yes', 'y', 'j', 'ok', 'x', 'true', 'go'):
        return True
    if s in ('nee', 'no', 'n', 'false', 'niet ok', 'not ok'):
        return False
    if s in ('', '-', 'y/n', 'j/n', 'ja/nee'):
        return None
    raise ValueError(v)


def to_time(v):
    if v is None:
        return None
    if isinstance(v, datetime.datetime):
        v = v.time()
    if isinstance(v, datetime.time):
        return v.strftime('%H:%M')
    if isinstance(v, (int, float)) and 0 <= v < 1:  # Excel day fraction
        m = round(v * 24 * 60)
        return f'{m // 60:02d}:{m % 60:02d}'
    if isinstance(v, (int, float)) and 0 <= v < 24:  # an hour
        return f'{int(v):02d}:{round((v % 1) * 60):02d}'
    s = str(v).strip().replace('.', ':')
    m = re.match(r'^(\d{1,2})(?::(\d{2}))?', s)
    if m and int(m.group(1)) < 24:
        return f'{int(m.group(1)):02d}:{m.group(2) or "00"}'
    raise ValueError(v)


def to_date(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    if v is None:
        return None
    s = str(v).strip()
    m = re.match(r'^(\d{4})-(\d{1,2})-(\d{1,2})', s)  # ISO date typed as text
    if m:
        return datetime.date(int(m.group(1)), int(m.group(2)), int(m.group(3))).isoformat()
    m = re.match(r'^(\d{1,2})[-/.](\d{1,2})[-/.](\d{2,4})$', s)
    if m:
        d, mo, y = (int(x) for x in m.groups())
        return datetime.date(y + 2000 if y < 100 else y, mo, d).isoformat()
    raise ValueError(v)


def to_timestamp(v, hour=None):
    """Local Excel date(time) -> ISO timestamp with offset."""
    if isinstance(v, datetime.datetime):
        dt = v
    elif isinstance(v, datetime.date):
        dt = datetime.datetime(v.year, v.month, v.day)
    elif v is None:
        return None
    else:
        dt = datetime.datetime.fromisoformat(to_date(v))
    if hour is not None:
        h, m = (int(x) for x in hour.split(':'))
        dt = dt.replace(hour=h, minute=m, second=0)
    return dt.replace(tzinfo=TZ).isoformat()


def people(v):
    """Initials in a cell ('TM/EP', 'BS EL', 'ns') -> ['TM', 'EP']."""
    if v is None:
        return []
    return [t.upper() for t in PEOPLE_SPLIT.split(str(v)) if t and t.strip()]


class Field:
    """Reads one cell into a typed value and records it in the review."""

    def __init__(self, book, sheet):
        self.b, self.sheet = book, sheet

    def __call__(self, ref, label, target, kind, allow_formula=False):
        v, status = self.b.get(self.sheet, ref, allow_formula)
        if status in ('formula', 'error') and v is None:
            raw = self.b.raw(self.sheet, ref)
            if raw is not None:
                self.b.note(self.sheet, ref, label, raw, target, 'not_imported', status)
            return None
        if v is None:
            return None
        try:
            out = {'num': to_num, 'bool': to_bool, 'time': to_time, 'date': to_date,
                   'text': lambda x: str(x).strip() or None, 'timestamp': to_timestamp,
                   'people': people}[kind](v)
        except (ValueError, TypeError):
            self.b.note(self.sheet, ref, label, v, target, 'suspicious', f'not a {kind}')
            self.b.warn(self.sheet, ref, f'{label}: "{v}" is not a {kind}')
            return None
        if out is None or out == []:
            return None
        self.b.note(self.sheet, ref, label, v, target, status)
        return out


# ------------------------------------------------------------------- Brewing

BREW_FIRST_COL = 'S'
BREW_COUNT = 9
# Dutch label in column P -> (batch_brews column, kind); None = calculated, skipped.
BREW_ROWS = {
    'order nr, software': ('software_order_number', 'text'),
    'start inmaischen (dd-mm uu:mm)': ('mash_in_started_at', 'timestamp'),
    'start inmaischen': ('mash_in_started_at', 'timestamp'),
    'maisch water': ('mash_water_hl', 'num'),
    'ph na inmaischen': ('mash_in_ph', 'num'),
    'correctie zuur': ('mash_acid_l', 'num'),
    'ph na correctie': ('mash_corrected_ph', 'num'),
    'joodnormaal test ok?': ('is_iodine_test_ok', 'bool'),
    'afwijking maische schema?': ('has_mash_schedule_deviation', 'bool'),
    'spoelafloop': ('sparge_runnings_hl', 'num'),
    'plato 1e wortvolume': ('first_wort_extract', 'num'),
    'extract laatste afloop': ('last_runnings_extract', 'num'),
    'eerste afloop': ('first_wort_volume_hl', 'num'),
    'cumm, sparging vol,': ('cumulative_sparge_volume_hl', 'num'),
    'last run ph': ('last_runnings_ph', 'num'),
    'filtratie tijd': ('lauter_time_min', 'num'),
    'ph begin kook': ('boil_start_ph', 'num'),
    'extract begin kook': ('boil_start_extract', 'num'),
    'v begin kook (80°c)': ('boil_start_volume_hl', 'num'),
    'zuurcorrectie': ('boil_acid_l', 'num'),
    'ph na zuurcorrectie / tijd': ('boil_acid_ph_note', 'text'),
    'totale kooktijd': ('boil_time_min', 'num'),
    'verdamping': None,
    'extra toevoeging': ('extra_addition', 'text'),
    'extract einde kook': ('boil_end_extract', 'num'),
    'volume whirlpool': ('whirlpool_volume_hl', 'num'),
    'werkelijke rust tijd': ('whirlpool_rest_min', 'num'),
    'beluchting': ('aeration_ppm', 'num'),
    'water additie': ('water_addition_hl', 'num'),
    'volume brouwsel': ('brew_volume_hl', 'num'),
    'volume glt': None,
    'extract tank': None,
    'hl*plato': None,
    'einde brouwen (dd-mm uu:mm)': ('brewing_ended_at', 'timestamp'),
    'einde brouwen': ('brewing_ended_at', 'timestamp'),
    'totale tijd brouwen': None,
}
# Labels that differ by section: "Extract" and "pH" under Whirlpool / end of boil.
BREW_SECTION_ROWS = {('whirlpool', 'extract'): ('whirlpool_extract', 'num'),
                     ('whirlpool', 'ph'): ('whirlpool_ph', 'num'),
                     # JK label of P19 ("Cumm, Sparging vol," at WP); the fixed
                     # "Spoelwater" row in Mashing is the rinse water.
                     ('filtratie & voorloopvat', 'spoelwater'): ('sparge_water_hl', 'num'),
                     ('maischen', 'spoelwater'): ('mash_rinse_water_hl', 'num')}
# Section header rows (no measured value; text per brew becomes a brew comment).
BREW_SECTIONS = {'maischen': 'Mashing', 'filtratie & voorloopvat': 'Filtration',
                 'koken': 'Boiling', 'whirlpool/ einde kook': 'Whirlpool'}
YEAST_ROWS = {
    'temp start': ('yeast_temp_start', 'num'), 'temp max': ('yeast_temp_max', 'num'),
    'giststam': ('yeast_strain', 'text'), 'gerehydrateerd?': ('yeast_rehydrated', 'bool'),
    'gewicht [gram]': ('yeast_weight', 'num'), 'bij 1e brouwsel': ('yeast_with_first_brew', 'bool'),
    'geënt op (datum)': ('yeast_pitched_on', 'date'), 'door': ('yeast_pitched_by', 'people'),
    'programma aan?': ('yeast_program_on', 'bool'),
}


def read_brewing(b, out):
    sh = 'Brouwen'
    if not b.has(sh):
        b.warn(sh, '-', 'sheet missing')
        return
    F = Field(b, sh)
    brews = {n: {'brew_number': n} for n in range(1, BREW_COUNT + 1)}
    comments = []
    section = None
    first = next((r for r, l in b.rows(sh, 'P') if l.startswith('brouwsel nr')), 3)
    for r, label in b.rows(sh, 'P'):
        if r <= first:
            continue  # title rows and the brew numbers
        if label in BREW_SECTIONS:
            section = label
        if label.startswith('whirlpool'):
            section = 'whirlpool'
        target = BREW_SECTION_ROWS.get((section, label))
        if target is None and label in BREW_ROWS:
            target = BREW_ROWS[label]
            if target is None:
                continue
        for n in range(1, BREW_COUNT + 1):
            vc = col(BREW_FIRST_COL, 2 * (n - 1))
            uc = col(vc, 1)
            ref = f'{vc}{r}'
            if label == 'naam brouwer':
                start = F(ref, f'Brewer start · brew {n}', 'batch_brews.brewer_id', 'people')
                end = F(f'{uc}{r}', f'Brewer end · brew {n}', 'batch_brews.brewer_end_id', 'people')
                if start:
                    brews[n]['brewer'] = start[0]
                    if len(start) > 1:
                        b.warn(sh, ref, f'brew {n}: more than one brewer start ({"/".join(start)}); first used')
                if end:
                    brews[n]['brewer_end'] = end[0]
                continue
            if label in BREW_SECTIONS or (target is None and label not in BREW_ROWS):
                # Header row or unknown label: text per brew is kept as a comment.
                for c in (vc, uc):
                    v, st = b.get(sh, f'{c}{r}')
                    if v is not None and st == 'read' and not (c == uc and norm(v) in UNITS):
                        name = BREW_SECTIONS.get(label, b.raw(sh, f'P{r}'))
                        comments.append({'brew_number': n, 'comment': f'{name}: {v}'})
                        b.note(sh, f'{c}{r}', f'{name} · brew {n}', v, 'batch_brew_comments',
                               'read' if label in BREW_SECTIONS else 'suspicious',
                               '' if label in BREW_SECTIONS else 'unknown row label; kept as comment')
                continue
            key, kind = target
            val = F(ref, f'{b.raw(sh, f"P{r}")} · brew {n}', f'batch_brews.{key}', kind)
            if val is None:
                v = b.raw(sh, ref)
                if isinstance(v, str) and v.strip() not in ('', '-') and not b.is_formula(sh, ref):
                    # Text in a number cell: keep it as a brew comment.
                    comments.append({'brew_number': n, 'comment': f'{b.raw(sh, f"P{r}")}: {v.strip()}'})
                continue
            brews[n][key] = val
            u, ust = b.get(sh, f'{uc}{r}')
            if u is not None and ust == 'read' and norm(u) not in UNITS:
                comments.append({'brew_number': n, 'comment': f'{b.raw(sh, f"P{r}")}: {u}'})
                b.note(sh, f'{uc}{r}', f'{b.raw(sh, f"P{r}")} · brew {n} (unit cell)', u,
                       'batch_brew_comments', 'suspicious', 'text in the unit cell; kept as comment')
    out['brews'] = [x for x in brews.values() if any(k != 'brew_number' for k in x)]
    # Remarks of the batch ("Opmerkingen :" in column A, below the recipe block).
    for r, label in b.rows(sh, 'A'):
        if label.startswith('opmerkingen'):
            x = r + 1
            while x <= b.v[sh].max_row and not norm(b.raw(sh, f'K{x}')).startswith('programma'):
                v, st = b.get(sh, f'A{x}')
                if v is not None and st == 'read' and str(v).strip() not in ('', '-', '0'):
                    comments.append({'brew_number': None, 'comment': str(v).strip()})
                    b.note(sh, f'A{x}', 'Remarks (Opmerkingen)', v, 'batch_brew_comments', 'read', 'all brews')
                x += 1
                if x > r + 15:
                    break
            break
    for r, label in b.rows(sh, 'G'):
        if label == 'glt':
            n = F(f'I{r}', 'GLT', 'batches.glt_tank_id', 'num')
            if n is not None:
                out['batch']['glt_tank'] = int(n)
    out['brew_comments'] = comments
    for r, label in b.rows(sh, 'K'):
        if label in YEAST_ROWS:
            key, kind = YEAST_ROWS[label]
            val = F(f'N{r}', b.raw(sh, f'K{r}'), f'batches.{key}', kind, allow_formula=True)
            if key == 'yeast_strain' and isinstance(val, str):
                val = val.rstrip(' -') or None  # typed with a trailing dash ("US-05-")
            if val is not None:
                out['batch'][key] = val[0] if kind == 'people' else val


UNITS = {'hl', '°p', '⁰p', 'min', 'liter', 'l', 'ppm', '%', 'uur', 'hl°p', 'kj/s', '°c', 'kg', 'g'}

# ------------------------------------------------------------ Fermentation

TAGS = [  # (text in the action cell, tag code); first match wins
    (r'24\s*h', '24h_purge'), (r'purge', 'purge'), (r'dry\s*warm', 'dry_warm'),
    (r'dry\s*hop', 'dryhopping'), (r'koel.*16|16.*koel|cooling.*16', 'cooling_16c'),
    (r'koel|cool', 'cooling'), (r'toevoeg|addition', 'addition'),
    (r'sample|monster|lab', 'sample_lab'), (r'samenvoeg|blend', 'blending'),
    (r'spund', 'spunding'), (r'proe|tast', 'tasting'), (r'overpomp|transfer', 'transfer'),
    (r'transport', 'transport'), (r'verwerk|process', 'processing'),
]


def tag_for(text):
    for rx, code in TAGS:
        if re.search(rx, text, re.I):
            return code
    return 'other'


def read_fermentation(b, out):
    sh = 'Gistkaart Invoer'
    if not b.has(sh):
        b.warn(sh, '-', 'sheet missing')
        return
    F = Field(b, sh)
    # GLT: typed on Brewing (I5, label "GLT" in G5); here it is a formula copy.
    if 'glt_tank' not in out['batch']:
        n = F('D4', 'GLT', 'batches.glt_tank_id', 'num', allow_formula=True)
        if n is not None:
            out['batch']['glt_tank'] = int(n)
    n = F('E4', 'LT', 'batches.lt_tank_id', 'num')
    if n is not None:
        out['batch']['lt_tank'] = int(n)
    # Header row with "Datum": columns by label (layout is the same in the samples).
    hdr = {norm(b.raw(sh, f'{get_column_letter(c)}12')): get_column_letter(c) for c in range(1, 24)
           if b.raw(sh, f'{get_column_letter(c)}12')}
    C = {k: hdr.get(k) for k in ('plato', 'temp hoog', 'temp laag', 'ph', 'datum', 'tijd', 'actie',
                                 'opmerkingen', 'operator')}
    if not C['datum']:
        b.warn(sh, '12', 'header row not found')
        return
    entries = []
    ws = b.v[sh]
    for r in range(13, ws.max_row + 1):
        date_v = b.raw(sh, f'{C["datum"]}{r}')
        if not isinstance(date_v, (datetime.datetime, datetime.date)):
            continue
        cells = {k: b.get(sh, f'{C[k]}{r}') for k in C if C[k] and k not in ('datum',)}
        filled = {k: v for k, (v, st) in cells.items() if v is not None and st == 'read'}
        if r == 13 or not filled or set(filled) == {'tijd'}:
            # Row 13 is Time 0 ("laatste brouwsel"); the app makes that entry itself.
            continue
        e = {}
        e['extract'] = F(f'{C["plato"]}{r}', f'Plato {date_v:%d-%m}', 'fermentation_log_entries.extract', 'num')
        e['temp_high'] = F(f'{C["temp hoog"]}{r}', f'Temp high {date_v:%d-%m}', 'fermentation_log_entries.temp_high', 'num')
        e['temp_low'] = F(f'{C["temp laag"]}{r}', f'Temp low {date_v:%d-%m}', 'fermentation_log_entries.temp_low', 'num')
        e['ph'] = F(f'{C["ph"]}{r}', f'pH {date_v:%d-%m}', 'fermentation_log_entries.ph', 'num')
        hour = F(f'{C["tijd"]}{r}', f'Hour {date_v:%d-%m}', 'fermentation_log_entries.measured_at', 'time')
        if hour is None:
            hour = '12:00'
            b.note(sh, f'{C["tijd"]}{r}', f'Hour {date_v:%d-%m}', None, 'fermentation_log_entries.measured_at',
                   'missing', 'no hour: 12:00 used')
        e['measured_at'] = to_timestamp(date_v, hour)
        comment = []
        action = F(f'{C["actie"]}{r}', f'Action {date_v:%d-%m}', 'fermentation_log_entry_tags', 'text')
        tags = []
        if action:
            code = tag_for(action)
            tags.append(code)
            if code == 'other':
                comment.append(action)
        text = F(f'{C["opmerkingen"]}{r}', f'Comment {date_v:%d-%m}', 'fermentation_log_entries.comment', 'text')
        if text:
            comment.append(text)
        ops = F(f'{C["operator"]}{r}', f'Operator {date_v:%d-%m}', 'fermentation_log_entries.created_by', 'people') or []
        if len(ops) > 1:
            comment.append('Op. ' + '/'.join(ops))
        e['comment'] = ' · '.join(comment) or None
        e['tags'] = tags
        e['operator'] = ops[0] if ops else None
        e['row'] = r
        entries.append({k: v for k, v in e.items() if v is not None or k in ('extract',)})
    out['fermentation_log'] = entries


# --------------------------------------------------------------- Processing

def read_processing(b, out):
    sh = 'Verwerking'
    if not b.has(sh):
        b.warn(sh, '-', 'sheet missing')
        return
    F = Field(b, sh)
    labels = b.rows(sh, 'A')
    starts = [r for r, l in labels if l.startswith('name operator start')]
    hbt_row = next((r for r, l in labels if l.startswith('bright beer tanks')), None)
    sessions, comments, tasting = [], [], []
    for i, r in enumerate(starts):
        end_row = starts[i + 1] if i + 1 < len(starts) else (hbt_row or b.v[sh].max_row)
        date = F(f'D{r + 2}', f'Date of processing · session {i + 1}', 'processing_sessions.processing_date', 'date')
        started = F(f'D{r}', f'Operator start · session {i + 1}', 'processing_sessions.started_by', 'people')
        if not date and not started:
            continue  # empty block (template defaults only)
        s = {'session_number': len(sessions) + 1, 'processing_date': date}
        ended = F(f'D{r + 1}', f'Operator end · session {i + 1}', 'processing_sessions.ended_by', 'people')
        s['started_by'] = started[0] if started else None
        s['ended_by'] = ended[0] if ended else None
        for who, lst in (('start', started), ('end', ended)):
            if lst and len(lst) > 1:
                comments.append({'comment_time': None,
                                 'comment': f'Session {s["session_number"]}: operators {who} {"/".join(lst)}'})
        s['start_time'] = F(f'P{r}', 'Start of processing', 'processing_sessions.start_time', 'time')
        s['end_time'] = F(f'P{r + 1}', 'End of processing', 'processing_sessions.end_time', 'time')
        block = {l: rr for rr, l in labels if r <= rr < end_row}
        rr = next((x for l, x in block.items() if l.startswith('processed from')), None)
        if rr:
            t = F(f'D{rr}', 'Processed from (G)LT', 'processing_sessions.tank_id', 'num')
            s['tank'] = int(t) if t is not None else None
            o = norm(b.raw(sh, f'G{rr}'))
            s['outlet'] = 'standpipe' if o.startswith('standp') else 'bottom_valve' if o.startswith('bodem') else None
            b.note(sh, f'G{rr}', 'Outlet', b.raw(sh, f'G{rr}'), 'processing_sessions.outlet', 'read' if s['outlet'] else 'missing')
        rr = next((x for l, x in block.items() if l.startswith('values in (g)lt')), None)
        if rr:
            s['centrifuge_in_o2_ppb'] = F(f'D{rr}', 'O2 in (G)LT before start', 'processing_sessions.centrifuge_in_o2_ppb', 'num')
            s['centrifuge_in_co2_gl'] = F(f'G{rr}', 'CO2 in (G)LT before start', 'processing_sessions.centrifuge_in_co2_gl', 'num')
            s['discharge_count'] = F(f'L{rr}', 'Losses (discharges)', 'processing_sessions.discharge_count', 'num')
            s['discharge_interval_min'] = F(f'O{rr}', 'Losses frequency (min)', 'processing_sessions.discharge_interval_min', 'num')
            if s['discharge_count'] is not None:
                s['discharge_count'] = int(s['discharge_count'])
        rr = next((x for l, x in block.items() if l.startswith('centrifuge start out')), None)
        if rr:
            s['centrifuge_out_o2_ppb'] = F(f'D{rr}', 'O2 centrifuge start out', 'processing_sessions.centrifuge_out_o2_ppb', 'num')
            s['centrifuge_out_co2_gl'] = F(f'G{rr}', 'CO2 centrifuge start out', 'processing_sessions.centrifuge_out_co2_gl', 'num')
            # Flow / hl / turbidity rows below the header: the first is the start.
            first = True
            for x in range(rr + 1, rr + 6):
                if norm(b.raw(sh, f'A{x}')):
                    break
                flow, turb = b.get(sh, f'J{x}')[0], b.get(sh, f'P{x}')[0]
                if flow is None and turb is None:
                    continue
                if first:
                    s['start_flow_hl_h'] = F(f'J{x}', 'Start flow', 'processing_sessions.start_flow_hl_h', 'num')
                    s['start_turbidity_ebc'] = F(f'P{x}', 'Start turbidity', 'processing_sessions.start_turbidity_ebc', 'num')
                    first = False
                else:
                    b.note(sh, f'J{x}', 'Centrifuge flow / hl / turbidity', f'{flow} / {b.raw(sh, f"M{x}")} / {turb}',
                           '-', 'not_imported', 'only the start values have a column')
        rr = next((x for l, x in block.items() if l.startswith('pasteurised with')), None)
        if rr:
            # Averages go onto every HBT (Jaap, 21-10); avg temp out has no column.
            pu = F(f'D{rr}', 'Pasteurised with (PU)', 'batch_hbts.avg_pu', 'num')
            turb = F(f'P{rr}', 'Avg. turbidity', 'batch_hbts.avg_turbidity_ebc', 'num')
            v, st = b.get(sh, f'J{rr}')
            if v is not None:
                b.note(sh, f'J{rr}', 'Avg. temp out', v, '-', 'not_imported', 'no column')
            s['_avg'] = {'pu': pu, 'turbidity': turb}
        rr = next((x for l, x in block.items() if l.startswith('time [hh:mm]')), None)
        rt = next((x for l, x in block.items() if l.startswith('comments tasting')), None)
        if rr:
            for x in range(rr + 1, rt or end_row):
                c = F(f'C{x}', 'Comment processing', 'processing_comments.comment', 'text')
                if c:
                    comments.append({'comment_time': F(f'A{x}', 'Comment time', 'processing_comments.comment_time', 'time'),
                                     'comment': c})
        if rt:
            text = F(f'C{rt}', 'Comments tasting', 'batches.tasting_comment', 'text')
            res = norm(b.raw(sh, f'R{rt}'))
            result = 'go' if res == 'go' else 'no_go' if res.replace(' ', '') in ('notgo', 'nogo', 'no-go') else None
            b.note(sh, f'R{rt}', 'Go or not go', b.raw(sh, f'R{rt}'), 'batches.tasting_result', 'read' if result else 'missing')
            tasting.append((s['session_number'], text, result))
        sessions.append({k: v for k, v in s.items() if v is not None})
    out['processing_sessions'] = sessions
    out['processing_comments'] = comments
    if tasting:
        texts = [f'Session {n}: {t}' if len(tasting) > 1 else t for n, t, _ in tasting if t]
        out['batch']['tasting_comment'] = ' · '.join(texts) or None
        results = [x for _, _, x in tasting if x]
        if results:
            out['batch']['tasting_result'] = results[-1]
    read_hbts(b, out, labels, hbt_row)
    for r, l in b.rows(sh, 'K'):
        if l.startswith('glt empty'):
            if not sessions and not out.get('transfers'):
                continue  # nothing processed: the cell is the template default
            v = F(f'R{r}', 'GLT empty?', 'batches.glt_empty', 'bool')
            if v is not None:
                out['batch']['glt_empty'] = v
    for r, l in labels:
        if l.startswith('measurement after pushing'):
            for c in ('D', 'F'):
                v, st = b.get(sh, f'{c}{r}')
                if v is not None:
                    b.note(sh, f'{c}{r}', 'Measurement after pushing over', v, '-', 'not_imported', 'no column')
    for r, l in b.rows(sh, 'H'):
        if l.startswith('start date lagering'):
            v, st = b.get(sh, f'N{r}')
            if v is not None:
                b.note(sh, f'N{r}', 'Start date lagering', v, '-', 'not_imported',
                       'the app derives lagering from the cooling tag in the log')


def parse_tank(v):
    """('hbt', number) for an HBT cell (a number or 'HBT 16'); ('other', 'LT51') for any other tank."""
    if v is None:
        return 'other', None
    try:
        n = to_num(v)
        if n is not None:
            return 'hbt', int(n)
    except ValueError:
        pass
    t = re.sub(r'\s+', '', str(v)).upper()
    m = re.match(r'^HBT(\d+)$', t)
    if m:
        return 'hbt', int(m.group(1))
    return 'other', t or None


def read_hbts(b, out, labels, hbt_row):
    sh = 'Verwerking'
    F = Field(b, sh)
    if not hbt_row:
        out['hbts'] = []
        return
    groups, cur = [], None
    for r, l in labels:
        if r <= hbt_row or l != 'hbt':
            if r > hbt_row and l.startswith('end value hbt'):
                break
            continue
        tank = b.get(sh, f'B{r}')[0]
        level = b.get(sh, f'C{r}')[0]
        if tank is None and level is None:
            cur = None  # an empty row closes a fill
            continue
        if tank is not None:
            kind, code = parse_tank(tank)
            if kind != 'hbt':
                b.warn(sh, f'B{r}', f'HBT "{tank}" is not an HBT: row not imported')
                continue
            tank = code
        else:
            tank = cur['tank'] if cur else None
        if cur is None or cur['tank'] != tank:
            cur = {'tank': tank, 'readings': [], 'formats': set()}
            groups.append(cur)
        target = norm(b.raw(sh, f'J{r}'))
        fmt = {'can': 'can', 'blik': 'can', 'keg': 'keg', 'fust': 'keg', 'bottle': 'bottle', 'fles': 'bottle'}.get(target)
        reading = {
            'sort_order': len(cur['readings']) + 1,
            'level_hl': F(f'C{r}', f'HBT{tank} level', 'hbt_readings.level_hl', 'num'),
            'o2_ppb': F(f'E{r}', f'HBT{tank} O2', 'hbt_readings.o2_ppb', 'num'),
            'co2_gl': F(f'G{r}', f'HBT{tank} CO2', 'hbt_readings.co2_gl', 'num'),
            'co2_target': fmt,
            'comment': F(f'M{r}', f'HBT{tank} comment', 'hbt_readings.comment', 'text'),
        }
        if fmt:
            cur['formats'].add(fmt)
        cur['readings'].append({k: v for k, v in reading.items() if v is not None})
    # End values, in the same order as the fills.
    ends = [r for r, l in labels if r > hbt_row and l.startswith('end value hbt')]
    used = 0
    for r in ends:
        tank = b.get(sh, f'D{r}')[0]
        level = F(f'E{r}', 'End value HBT', 'hbt_readings.level_hl (last)', 'num')
        if tank is None and level is None:
            continue
        note = ' '.join(str(b.raw(sh, f'{c}{r}') or '') for c in ('K', 'L', 'M', 'N')).strip()
        m = re.search(r'(?<!\d)(2\d{4})(?!\d)', note)
        kind, code = parse_tank(tank)
        if kind != 'hbt':
            # Not an HBT: the beer went elsewhere from the batch's own tank.
            if m and level:
                out.setdefault('transfers', []).append({
                    'to_batch_number': m.group(1), 'volume_hl': level,
                    'comment': f'via {code}: {note}' if code else note,
                    'from_tank': out['batch'].get('lt_tank') or out['batch'].get('glt_tank'),
                    'transferred_at': None, 'source': f'{sh}!E{r}'})
                b.note(sh, f'E{r}', f'End value {code or "(no tank)"}', f'{level} ({note})', 'batch_transfers', 'suspicious',
                       f'transfer to batch {m.group(1)}: check date and volume')
            elif level:
                out.setdefault('other_end_values', []).append({
                    'kind': 'end', 'tank': code, 'volume_hl': level, 'comment': note, 'cell': f'{sh}!E{r}',
                    'o2_ppb': b.get(sh, f'G{r}')[0], 'co2_gl': b.get(sh, f'I{r}')[0]})
            continue
        tank = code
        g = next((x for x in groups[used:] if x['tank'] == tank and 'end' not in x), None)
        if g is None:
            b.warn(sh, f'D{r}', f'end value for HBT{tank} without readings: new fill made')
            g = {'tank': tank, 'readings': [], 'formats': set()}
            groups.append(g)
        g['end'] = level
        if m and level:
            # This fill went to another batch (a split from the HBT).
            g['split_to'] = m.group(1)
            out.setdefault('transfers', []).append({
                'to_batch_number': m.group(1), 'volume_hl': level, 'comment': note,
                'from_tank': tank, 'from_hbt': True, 'transferred_at': None, 'source': f'{sh}!E{r}'})
            b.note(sh, f'K{r}', f'HBT{tank} end value to a batch', f'{level} ({note})', 'batch_transfers', 'suspicious',
                   f'split from HBT{tank} to {m.group(1)}: check date and volume')
        last = g['readings'][-1] if g['readings'] else {}
        end_reading = {'level_hl': level, 'o2_ppb': F(f'G{r}', 'End value O2', 'hbt_readings.o2_ppb', 'num'),
                       'co2_gl': F(f'I{r}', 'End value CO2', 'hbt_readings.co2_gl', 'num')}
        if any(last.get(k) != v for k, v in end_reading.items() if v is not None):
            end_reading['sort_order'] = len(g['readings']) + 1
            end_reading['comment'] = 'End value'
            if last.get('co2_target'):
                end_reading['co2_target'] = last['co2_target']
            g['readings'].append({k: v for k, v in end_reading.items() if v is not None})
        used = groups.index(g) + 1
    for g in groups:
        kept = [r for r in g['readings'] if any(r.get(k) is not None for k in ('level_hl', 'o2_ppb', 'co2_gl', 'comment'))]
        for i, r in enumerate(kept):
            r['sort_order'] = i + 1
        g['readings'] = kept
    groups = [g for g in groups if g['readings'] or g.get('end')]
    sessions = out.get('processing_sessions', [])
    hbts = []
    for i, g in enumerate(groups):
        h = {'tank': g['tank'], 'sort_order': i + 1, 'intended_formats': sorted(g['formats']),
             'readings': g['readings'], 'end_level_hl': g.get('end')}
        if g.get('split_to'):
            h['split_to'] = g['split_to']
        if len(sessions) == 1:
            h['processing_session_number'] = 1  # one session: every fill comes from it
        for key, col_ in (('pu', 'avg_pu'), ('turbidity', 'avg_turbidity_ebc')):
            # The same averages on every HBT: the mean over the sessions (Jaap, 21-10).
            vals = [x['_avg'][key] for x in sessions if x.get('_avg', {}).get(key) is not None]
            if vals:
                h[col_] = round(sum(vals) / len(vals), 2)
        hbts.append(h)
    if len(sessions) > 1 and groups:
        b.warn(sh, f'A{hbt_row}', f'{len(sessions)} sessions: HBT readings are not linked to a session (not in the workbook)')
    for s in sessions:
        s.pop('_avg', None)
    out['hbts'] = hbts


# ------------------------------------------------------------------ Filling

FILL_COLS = ['M', 'N', 'O', 'P', 'Q', 'R']       # Afvullen
REPORT_COLS = ['D', 'E', 'F', 'G', 'H', 'I']     # Afvulverslag, same order
MEAS_ROWS = [  # label prefix in column I -> filling_measurements column
    ('time', 'measured_time', 'time'), ('tpo', 'tpo_ppb', 'num'), ('hso', 'hso_ppb', 'num'),
    ('do ', 'do_ppb', 'num'), ('hs vol', 'headspace_volume', 'num'), ('druk', 'pressure_bar', 'num'),
    ('temp', 'temp_c', 'num'), ('co2 fles', 'co2_gl', 'num'), ('co2 fust', 'keg_co2_gl', 'num'),
    ('zuurstof fust', 'keg_o2_ppb', 'num'),
]
ARTICLES = [  # Dutch name prefix in Afvulverslag column B -> packaging_articles.code
    ('doos 24x33cl - blanco', 'box_24x33_blank'), ('doos 24x33cl blik', 'box_24x33_can'),
    ('doos 24x33cl', 'box_24x33'), ('doos 12x33cl - blanco', 'box_12x33_blank'),
    ('doos 12x33cl blik', 'box_12x33_can'), ('doos 12x33cl', 'box_12x33'),
    ('doos 6x75cl - blanco', 'box_6x75_blank'), ('doos 6x75cl', 'box_6x75'),
    ('krat 12x75cl', 'crate_12x75_blank'), ('doos 12x44cl blik', 'box_12x44_can'),
    ('doos 24x44cl blik', 'box_24x44_can'), ('20l keykeg', 'keykeg_20'), ('30l keykeg', 'keykeg_30'),
    ('20l jopen/ku fust', 'jopen_keg_20'), ('30l jopen fust', 'jopen_keg_30'), ('50l jopen fust', 'jopen_keg_50'),
    ('20l gbr, door vrouwen', 'gbr_door_vrouwen_keg_20'), ('20l global keg', None),
]
MATERIALS = [  # -> packaging_materials.code (None = no longer used: reported)
    ('fles 33cl', 'bottle_33'), ('fles 75cl', 'bottle_75'),
    ('stickerlabel 33cl blik', 'label_33_can'), ('stickerlabel 44cl blik', 'label_44_can'),
    ('stickerlabel 33cl - main', 'label_33_main'), ('stickerlabel 75cl - main', 'label_75_main'),
    ('stickerlabel 33cl - neck', None), ('stickerlabel 75cl - neck', None),
    ('omdozen 24x33cl blik', 'box_24x33_can'), ('omdozen 12x33cl blik', 'box_12x33_can'),
    ('omdozen 24x44cl blik', 'box_24x44_can'), ('omdozen 12x44cl blik', 'box_12x44_can'),
    ('omdozen 24x33cl - bedrukt', 'box_24x33_printed'), ('omdozen 24x33cl - blanco', 'box_24x33_blank'),
    ('omdozen 12x33cl - bedrukt', 'box_12x33_printed'), ('omdozen 12x33cl - blanco', 'box_12x33_blank'),
    ('omdozen 6x75cl - bedrukt', 'box_6x75_printed'), ('omdozen 6x75cl - blanco', 'box_6x75_blank'),
    ('kroonkurk', 'crown_cork'), ('kurk', 'cork'), ('musselet', 'muselet'),
    ('blik 33cl', 'can_33'), ('blik 44cl', 'can_44'), ('blik deksel', 'can_lid'),
    ('global keg 20l', None), ('jopen fust 20l', 'jopen_keg_20'), ('jopen fust 30l', 'jopen_keg_30'),
    ('jopen fust 50l', 'jopen_keg_50'), ('gdv 20l', 'gbr_door_vrouwen_keg_20'), ('fustdop', 'keg_cap'), ('keykegfustdop', 'keykeg_cap'),
    ('keykeg 20l', 'keykeg_20'), ('keykeg 30l', 'keykeg_30'),
]


def match(name, table):
    for prefix, code in table:
        if name.startswith(prefix):
            return True, code
    return False, None


def read_filling(b, out, ref):
    sh = 'Afvullen'
    if not b.has(sh):
        b.warn(sh, '-', 'sheet missing')
        return
    F = Field(b, sh)
    labels = dict((l, r) for r, l in reversed(b.rows(sh, 'I')))
    row = lambda prefix: next((r for l, r in labels.items() if l.startswith(prefix)), None)
    r_filled, r_sampled, r_check = row('afgevuld door'), row('monstername door'), row('controle 1e verpakking')
    r_tube1, r_tube2, r_burst = row('vulpijp-detectiecontrole 1e'), row('vulpijp-detectiecontrole ein'), row('klapfles beheerst')
    meas_starts = [r for r, l in b.rows(sh, 'I') if l == 'time']
    sessions = []
    for i, c in enumerate(FILL_COLS):
        date = F(f'{c}6', f'Filling date · column {c}', 'filling_sessions.filling_date', 'date')
        if not date:
            continue
        s = {'session_number': len(sessions) + 1, 'filling_date': date, 'column': i}
        hbt = F(f'{c}8', 'HBT nr', 'filling_sessions.batch_hbt_id', 'num')
        s['hbt_tank'] = int(hbt) if hbt is not None else None
        s['start_time'] = F(f'{c}10', 'Start filling', 'filling_sessions.start_time', 'time')
        s['end_time'] = F(f'{c}12', 'End filling', 'filling_sessions.end_time', 'time')
        filled = F(f'{c}{r_filled}', 'Filled by', 'filling_sessions.created_by', 'people') if r_filled else None
        extra = F(f'{c}{r_filled + 1}', 'Filled by (2nd line)', 'filling_sessions.created_by', 'people') if r_filled else None
        filled = (filled or []) + (extra or [])
        s['filled_by'] = filled[0] if filled else None
        if len(filled) > 1:
            s['also_filled_by'] = filled[1:]
        sampled = F(f'{c}{r_sampled}', 'Sampled by', 'filling_sessions.sampled_by', 'people') if r_sampled else None
        s['sampled_by'] = sampled[0] if sampled else None
        if r_check:
            chk = F(f'{c}{r_check}', 'First package check (two operators)', 'filling_sessions.is_first_package_checked', 'people')
            s['is_first_package_checked'] = True if chk else None
            if chk:
                s['first_package_checked_by'] = chk
        for r, key in ((r_tube1, 'is_fill_tube_checked_first'), (r_tube2, 'is_fill_tube_checked_end'),
                       (r_burst, 'are_burst_bottles_controlled')):
            if r:
                v, st = b.get(sh, f'{c}{r}')
                if v is not None and st == 'read':
                    try:
                        s[key] = to_bool(v) if norm(v) not in people_set(ref) else True
                    except ValueError:
                        s[key] = True  # initials or a remark: done
                    b.note(sh, f'{c}{r}', key, v, f'filling_sessions.{key}', 'read')
        meas = []
        for start in meas_starts:
            m = {}
            for off, (prefix, key, kind) in enumerate(MEAS_ROWS):
                lab = norm(b.raw(sh, f'I{start + off}'))
                if not lab.startswith(prefix.strip()):
                    b.warn(sh, f'I{start + off}', f'expected "{prefix}", found "{lab}"')
                    break
                v = F(f'{c}{start + off}', f'{b.raw(sh, f"I{start + off}")}', f'filling_measurements.{key}', kind)
                if v is not None:
                    m[key] = v
            if any(k != 'measured_time' for k in m):
                meas.append(m)
        s['measurements'] = meas
        sessions.append({k: v for k, v in s.items() if v is not None})
    read_filling_report(b, out, sessions, ref)
    flag_outliers(b, sessions)
    out['filling_sessions'] = sessions


def people_set(ref):
    return {u['initials'].lower() for u in ref['users'] if u.get('initials')}


def read_filling_report(b, out, sessions, ref):
    sh = 'Afvulverslag'
    if not b.has(sh):
        b.warn(sh, '-', 'sheet missing')
        return
    F = Field(b, sh)
    by_col = {s['column']: s for s in sessions}
    fmt = {a['code']: a['fmt'] for a in ref['articles']}
    for i, c in enumerate(REPORT_COLS):
        date = b.get(sh, f'{c}7')[0]
        if date is None:
            continue
        s = by_col.get(i)
        if s is None:
            # Old reports: a session written only on the filling report.
            hbt = b.get(sh, f'{c}9')[0]
            kind, code = parse_tank(hbt)
            if hbt is not None and kind != 'hbt':
                # Not from an HBT (e.g. "TT60" / "Kerk": tank beer for the Jopenkerk).
                vol = to_num(b.get(sh, f'{c}11')[0])
                out.setdefault('other_end_values', []).append({
                    'kind': 'report', 'tank': code, 'volume_hl': vol / 100 if vol and vol > 300 else vol,  # liters or hl
                    'date': to_date(date), 'comment': f'filling report column {c}', 'cell': f'{sh}!{c}9'})
                continue
            s = {'session_number': len(sessions) + 1, 'filling_date': to_date(date), 'column': i,
                 'hbt_tank': int(to_num(hbt)) if to_num(hbt) is not None else None, 'measurements': [],
                 'from_report_only': True}
            who = people(b.get(sh, f'{c}13')[0])
            if who:
                s['filled_by'] = who[0]
            sessions.append(s)
            by_col[i] = s
            b.note(sh, f'{c}7', 'Filling session (filling report only)', date, 'filling_sessions', 'suspicious',
                   'not on Afvullen: made from the filling report (date, HBT, volume, articles)')
        if to_date(date) != s['filling_date']:
            b.warn(sh, f'{c}7', f'date {to_date(date)} differs from Afvullen ({s["filling_date"]})')
        vol = F(f'{c}11', 'Volume (HL)', 'filling_sessions.volume_hl', 'num')
        if vol is not None:
            s['volume_hl'] = vol
        s['articles'], s['materials'] = [], []
    section = None
    for r, label in b.rows(sh, 'B'):
        if label.startswith('afvul ') or label.startswith('opboeken'):
            section = 'articles'
            continue
        if label.startswith('verbruiken'):
            section = 'materials'
            continue
        if label.startswith('gebruikt tbv samenvoegen'):
            # The old way of a split: liters booked off "TBV BATCH <number>".
            section = 'skip'
            for x in range(r + 1, r + 4):
                name = ' '.join(str(b.raw(sh, f'{c}{x}') or '') for c in ('B', 'C')).strip()
                m = re.search(r'\b(2\d{4})\b', name)
                for i, c in enumerate(REPORT_COLS):
                    v, st = b.get(sh, f'{c}{x}')
                    if v in (None, 0) or st != 'read':
                        continue
                    liters = to_num(v)
                    if not m:
                        date = b.get(sh, f'{c}7')[0]
                        out.setdefault('other_end_values', []).append({
                            'kind': 'tbv', 'tank': 'TT60' if re.search(r'tt\s*60|kerk', name, re.I) else None,
                            'volume_hl': round(liters / 100, 2), 'date': to_date(date) if date else None,
                            'comment': name, 'cell': f'{sh}!{c}{x}'})
                        continue
                    date = b.get(sh, f'{c}7')[0]
                    out.setdefault('report_transfers', []).append({
                        'to_batch_number': m.group(1), 'volume_hl': round(liters / 100, 2),
                        'date': to_date(date) if date else None, 'source': f'{sh}!{c}{x}'})
                    b.note(sh, f'{c}{x}', name, v, 'batch_transfers.volume_hl', 'read', 'liters / 100')
            continue
        if label.startswith('ingevoerd in mdrank'):
            section = 'skip'
            continue
        if section not in ('articles', 'materials'):
            continue
        known, code = match(label, ARTICLES if section == 'articles' else MATERIALS)
        for i, c in enumerate(REPORT_COLS):
            s = by_col.get(i)
            v, st = b.get(sh, f'{c}{r}')
            if v is None or st != 'read':
                continue
            q = F(f'{c}{r}', b.raw(sh, f'B{r}'), f'filling_session_{section}', 'num')
            if q is None or s is None or 'articles' not in s:
                continue
            if not known or code is None:
                b.note(sh, f'{c}{r}', b.raw(sh, f'B{r}'), v, f'filling_session_{section}', 'not_imported',
                       'no longer used (removed from Settings)' if known else 'unknown name')
                b.warn(sh, f'{c}{r}', f'"{b.raw(sh, f"B{r}")}" booked ({q:g}) but {"no longer used" if known else "unknown"}')
                continue
            s[section].append({'code': code, 'quantity': q})
    for s in sessions:
        formats = {fmt.get(a['code']) for a in s.get('articles', []) if a['quantity']}
        keg_meas = any('keg_co2_gl' in m or 'keg_o2_ppb' in m for m in s.get('measurements', []))
        if len(formats) == 1:
            s['package_format'] = formats.pop()
        elif not formats and keg_meas:
            s['package_format'] = 'keg'
        elif len(formats) > 1:
            b.warn('Afvulverslag', '-', f'session {s["session_number"]}: more than one package ({", ".join(sorted(formats))}); none set')


def flag_outliers(b, sessions):
    limits = {'tpo_ppb': 100, 'hso_ppb': 50, 'do_ppb': 50, 'keg_o2_ppb': 50}
    for s in sessions:
        for m in s.get('measurements', []):
            for k, lim in limits.items():
                if m.get(k) is not None and m[k] > 3 * lim:
                    b.warn('Afvullen', '-', f'session {s["session_number"]} {m.get("measured_time", "")}: '
                                           f'{k} {m[k]:g} is more than 3x the limit {lim} (check)')


# ------------------------------------------------- links between the parts

TANK_BEER = 'TT60'   # the transport tank for the Jopenkerk (package format "tank", Jaap 21-10)


def tank_beer(b, out):
    """Beer processed into TT60 for the Jopenkerk: a fill of TT60 plus one filling
    session of package "tank" (article tank_jopenkerk, counted in hl). The date
    comes from the filling report entry with that volume (a TT60 / "Kerk"
    column or a "TBV ..." row), else the last processing date. Everything else
    outside an HBT is reported for Jaap, not imported."""
    other = out.pop('other_end_values', [])
    used = set()
    for i, e in enumerate(other):
        if e['kind'] != 'end' or not (e['tank'] or '').startswith('TT'):
            continue
        used.add(i)
        vol = e['volume_hl']
        match = next((j for j, x in enumerate(other) if j not in used and x['kind'] in ('report', 'tbv')
                      and x.get('volume_hl') and abs(x['volume_hl'] - vol) <= max(0.5, 0.03 * vol)), None)
        if match is not None:
            used.add(match)
            date = other[match].get('date')
            # The same delivery written twice on the filling report (a column and a "TBV ..." row).
            for j, x in enumerate(other):
                if j not in used and x['kind'] in ('report', 'tbv') and x.get('tank') == TANK_BEER and x.get('volume_hl') \
                        and abs(x['volume_hl'] - vol) <= max(0.5, 0.03 * vol):
                    used.add(j)
                    date = date or x.get('date')
        else:
            date = None
        if not date:
            dates = [s['processing_date'] for s in out.get('processing_sessions', []) if s.get('processing_date')]
            date = max(dates) if dates else None
            b.warn('-', '-', f'tank beer {vol:g} hl ({e["tank"]}): '
                             + ('the filling report entry has no date' if match is not None else 'not on the filling report')
                             + f', date = last processing date {date}')
        reading = {'sort_order': 1, 'level_hl': vol, 'comment': 'End value (tank beer Jopenkerk)', 'co2_target': 'tank'}
        for k in ('o2_ppb', 'co2_gl'):
            if to_num(e.get(k)) is not None:
                reading[k] = to_num(e[k])
        hbts = out.setdefault('hbts', [])
        hbts.append({'tank': TANK_BEER, 'sort_order': len(hbts) + 1, 'intended_formats': ['tank'],
                     'readings': [reading], 'end_level_hl': vol})
        sessions = out.setdefault('filling_sessions', [])
        if not date:
            b.warn('-', '-', f'tank beer {vol:g} hl: no date at all, no filling session made')
            continue
        sessions.append({'session_number': len(sessions) + 1, 'filling_date': date, 'hbt_tank': TANK_BEER,
                         'package_format': 'tank', 'volume_hl': vol, 'measurements': [],
                         'articles': [{'code': 'tank_jopenkerk', 'quantity': vol}], 'materials': []})
        b.note(*e['cell'].split('!'), f'End value {e["tank"]}', f'{vol} ({e["comment"]})', 'filling_sessions (tank)', 'read',
               f'tank beer Jopenkerk, filled {date}' + (f' ({other[match]["cell"]})' if match is not None else ''))
    for i, e in enumerate(other):
        if i in used:
            continue
        sh, cell = e['cell'].split('!')
        b.note(sh, cell, f'{e["kind"]}: {e["tank"] or e["comment"]}', f'{e["volume_hl"]} hl ({e["comment"]})', '-',
               'not_imported', 'not from an HBT and not to a batch: ask Jaap')
        b.warn(sh, cell, f'{e["volume_hl"]:g} hl {e["tank"] or ""} ({e["comment"] or "no comment"}): not imported, ask Jaap')
        out.setdefault('other_end_values', []).append(e)


def apply_overrides(b, out, path=os.path.join(os.path.dirname(os.path.abspath(__file__)), 'overrides.json')):
    """One-off corrections agreed with Jaap (overrides.json), per batch number."""
    if not os.path.exists(path):
        return
    o = json.load(open(path)).get(out['batch_number'] or '')
    if not o:
        return
    for h in o.get('hbts', []):
        hbts = out.setdefault('hbts', [])
        hbts.append({**h, 'sort_order': len(hbts) + 1})
    for s in o.get('filling_sessions', []):
        sessions = out.setdefault('filling_sessions', [])
        sessions.append({**s, 'session_number': len(sessions) + 1})
    if o.get('drop_other_end_values'):
        out.pop('other_end_values', None)
        b.warnings[:] = [w for w in b.warnings if 'ask Jaap' not in w]
    b.note('-', '-', 'One-off correction', o.get('note', ''), 'overrides.json', 'read', 'agreed with Jaap')


def link_hbts(b, out):
    """Filling sessions -> HBT fill (index), and HBT finished.

    Sessions of the same tank go to its fills in order: a fill is finished, and
    the next session of that tank goes to the next fill, once the filled volume
    reaches its end value (within 3 %)."""
    hbts = out.get('hbts', [])
    by_tank = {}
    for i, h in enumerate(hbts):
        by_tank.setdefault(h['tank'], []).append(i)
    pointer, filled = {}, {}
    for s in sorted(out.get('filling_sessions', []), key=lambda x: (x['filling_date'], x['session_number'])):
        t = s.get('hbt_tank')
        if t not in by_tank:
            if t is not None:
                b.warn('Afvullen', '-', f'session {s["session_number"]}: HBT{t} has no readings in Processing')
            continue
        p = pointer.get(t, 0)
        if p >= len(by_tank[t]):
            p = len(by_tank[t]) - 1
            b.warn('Afvullen', '-', f'session {s["session_number"]}: HBT{t} filled after its last fill was finished')
        idx = by_tank[t][p]
        s['hbt_index'] = idx
        filled[idx] = filled.get(idx, 0) + (s.get('volume_hl') or 0)
        end = hbts[idx].get('end_level_hl')
        if end and filled[idx] >= end * 0.97:
            s['is_hbt_finished'] = True
            pointer[t] = p + 1
    for i, h in enumerate(hbts):
        end = h.get('end_level_hl')
        if h.get('split_to') and i in filled:
            b.warn('Afvullen', '-', f'HBT{h["tank"]} (fill {h["sort_order"]}) went to {h["split_to"]} but also has filling sessions')
        if end and i in filled and filled[i] < end * 0.97:
            b.warn('Afvullen', '-', f'HBT{h["tank"]} (fill {h["sort_order"]}): filled {filled[i]:g} of {end:g} hl, not finished')
        if not h['intended_formats']:
            fm = {s.get('package_format') for s in out.get('filling_sessions', []) if s.get('hbt_index') == i}
            h['intended_formats'] = sorted(x for x in fm if x)


def merge_transfers(b, out):
    """Transfers from the processing form ("End value" rows) and from the filling
    report ("TBV ... <batch>"), compared per receiving batch. The processing form
    can split one transfer over several HBTs (26093: 10 + 30 hl = 4000 l)."""
    ts = out.setdefault('transfers', [])
    report = {}
    for rt in out.pop('report_transfers', []):
        x = report.setdefault(rt['to_batch_number'], {'volume_hl': 0, 'date': rt['date'], 'source': []})
        x['volume_hl'] += rt['volume_hl']
        x['source'].append(rt['source'])
    for to, rt in report.items():
        mine = [t for t in ts if t['to_batch_number'] == to]
        if not mine:
            ts.append({'to_batch_number': to, 'volume_hl': round(rt['volume_hl'], 2), 'comment': None,
                       'from_tank': out['batch'].get('lt_tank') or out['batch'].get('glt_tank'),
                       'transferred_at': to_timestamp(rt['date'], '12:00') if rt['date'] else None,
                       'source': ' + '.join(rt['source'])})
            continue
        total = sum(t['volume_hl'] for t in mine)
        for t in mine:
            t['source'] += ' + ' + ' + '.join(rt['source'])
        if abs(total - rt['volume_hl']) > 0.05:
            b.warn('-', '-', f'transfer to {to}: {total:g} hl on Processing, {rt["volume_hl"]:g} hl on the filling report')
    if not ts:
        out.pop('transfers')


def date_transfers(b, out):
    for t in out.get('transfers', []):
        if t.get('transferred_at'):
            continue
        vol = f'{t["volume_hl"]:g}'
        pats = {vol, vol.replace('.', ',')}
        hits = [e for e in out.get('fermentation_log', []) if e.get('comment') and any(p in e['comment'] for p in pats)]
        if hits:
            t['transferred_at'] = hits[-1]['measured_at']
            t['date_from'] = f'log entry: {hits[-1]["comment"][:80]}'
            if 'transfer' not in hits[-1]['tags']:
                hits[-1]['tags'].append('transfer')
        else:
            b.warn('-', '-', f'transfer to {t["to_batch_number"]}: no date found (log has no entry with {vol} hl)')
        if t['from_tank'] is None:
            b.warn('-', '-', f'transfer to {t["to_batch_number"]}: from tank unknown')


def check_dates(b, out):
    starts = [x.get('mash_in_started_at') for x in out.get('brews', []) if x.get('mash_in_started_at')]
    if not starts:
        return
    first = min(starts)[:10]
    for e in out.get('fermentation_log', []):
        if e['measured_at'][:10] < first:
            b.warn('Gistkaart Invoer', f'J{e["row"]}', f'log entry {e["measured_at"][:10]} before the first brew {first}')
    for s in out.get('processing_sessions', []) + out.get('filling_sessions', []):
        d = s.get('processing_date') or s.get('filling_date')
        if d and d < first:
            b.warn('-', '-', f'date {d} before the first brew {first}')


def resolve(b, out, ref):
    """Tank numbers -> codes; check initials, tags, codes against the app."""
    tanks = ref['tanks']
    def tank_code(n, types):
        if isinstance(n, str):
            hits = [t['code'] for t in tanks if t['code'] == n]
            if hits:
                return hits[0]
        hits = [t['code'] for t in tanks if t['n'] == n and t['type'] in types]
        if len(hits) == 1:
            return hits[0]
        b.warn('-', '-', f'tank number {n} ({"/".join(types)}): {"not found" if not hits else "ambiguous " + ", ".join(hits)}')
        return None
    bt = out['batch']
    for key in ('glt_tank', 'lt_tank'):
        if key in bt:
            bt[key] = tank_code(bt[key], ('glt', 'ckt', 'lt'))
    for s in out.get('processing_sessions', []):
        if 'tank' in s:
            s['tank'] = tank_code(s['tank'], ('glt', 'ckt', 'lt'))
    for t in out.get('transfers', []):
        if isinstance(t.get('from_tank'), int):
            t['from_tank'] = tank_code(t['from_tank'], ('hbt',) if t.get('from_hbt') else ('glt', 'ckt', 'lt'))
    for h in out.get('hbts', []):
        h['tank'] = tank_code(h['tank'], ('hbt',))
    users = {u['initials'].upper() for u in ref['users'] if u.get('initials')}
    seen = set()
    def check(ini, where):
        if ini and ini not in users and (ini, where) not in seen:
            seen.add((ini, where))
            b.warn('-', '-', f'initials {ini} ({where}) are not a user')
    check(bt.get('yeast_pitched_by'), 'yeast pitched by')
    for x in out.get('brews', []):
        check(x.get('brewer'), 'brewer'); check(x.get('brewer_end'), 'brewer end')
    for e in out.get('fermentation_log', []):
        check(e.get('operator'), 'fermentation log')
    for s in out.get('processing_sessions', []):
        check(s.get('started_by'), 'processing'); check(s.get('ended_by'), 'processing')
    for s in out.get('filling_sessions', []):
        check(s.get('filled_by'), 'filling'); check(s.get('sampled_by'), 'filling')


def workbook_type(b):
    """generated / old (batch report) or barrel_aging; None = unknown layout."""
    names = {n.strip() for n in b.v}
    if 'Hout Rijping' in names:
        return 'barrel_aging'
    if {'Brouwen', 'Gistkaart Invoer', 'Verwerking', 'Afvullen', 'Afvulverslag'} <= names:
        return 'old' if 'EBU Berekening' in names else 'generated'
    return None


def read_workbook(path, ref):
    b = Book(path)
    name = os.path.basename(path)
    out = {'source': name, 'batch_number': None, 'batch': {}, 'type': workbook_type(b)}
    if out['type'] not in ('generated', 'old'):
        # Not a batch report this reader knows: held back, never guessed.
        m = re.match(r'^(\d{5})', name)
        out['batch_number'] = m.group(1) if m else None
        out['hold'] = ('barrel aging workbook: separate import (see import-2026.md)' if out['type'] == 'barrel_aging'
                       else 'unknown layout: sheets ' + ', '.join(b.v))
        out['review'], out['warnings'] = [], [out['hold']]
        return out
    for sh, cell in (('Brouwen', 'N3'), ('Verwerking', 'L3'), ('Afvullen', 'E3'), ('Afvulverslag', 'H3')):
        v = b.raw(sh, cell) if b.has(sh) else None
        if isinstance(v, (int, float)) and 20000 <= v < 30000 or isinstance(v, str) and re.fullmatch(r'\s*2\d{4}\s*', v):
            out['batch_number'] = str(int(to_num(v)))
            break
        if v is not None:
            b.warn(sh, cell, f'batch number cell holds "{v}": taken from the next sheet')
    m = re.match(r'^(\d{5})\s+(.*?)(?:\s+v(\d+)(?:\.(\d+))?)?(?:\s+(?:WP|JK))?\.xlsx?$', name, re.I)
    if m:
        ver = f'{m.group(3)}.{m.group(4) or 0}' if m.group(3) else None
        out['workbook_recipe'] = {'name': m.group(2), 'version': ver}
        if out['batch_number'] and out['batch_number'] != m.group(1):
            b.warn('-', '-', f'file name {m.group(1)} differs from the batch number {out["batch_number"]} in the workbook')
    read_brewing(b, out)
    read_fermentation(b, out)
    read_processing(b, out)
    read_filling(b, out, ref)
    tank_beer(b, out)
    apply_overrides(b, out)
    link_hbts(b, out)
    merge_transfers(b, out)
    date_transfers(b, out)
    check_dates(b, out)
    resolve(b, out, ref)
    for s in out.get('filling_sessions', []):
        s.pop('column', None)
    out['review'] = b.review
    out['warnings'] = b.warnings
    return out


def main():
    out_dir, ref_path, books = sys.argv[1], sys.argv[2], sys.argv[3:]
    ref = json.load(open(ref_path))
    os.makedirs(out_dir, exist_ok=True)
    failed = []
    for p in books:
        try:
            data = read_workbook(p, ref)
        except Exception as e:  # one bad workbook must not stop the run
            import traceback
            failed.append(os.path.basename(p))
            print(f'FAILED {os.path.basename(p)}: {type(e).__name__}: {e}')
            with open(os.path.join(out_dir, f'FAILED_{os.path.basename(p)}.txt'), 'w') as f:
                f.write(traceback.format_exc())
            continue
        dest = os.path.join(out_dir, f'{data["batch_number"] or os.path.basename(p)}.json')
        if os.path.exists(dest) and json.load(open(dest)).get('source') != data['source']:
            data['warnings'].append(f'batch number also in {json.load(open(dest))["source"]}: check')
            dest = dest[:-5] + '_2.json'
        with open(dest, 'w') as f:
            json.dump(data, f, indent=1, ensure_ascii=False, default=str)
        counts = {k: len(data.get(k, [])) for k in ('brews', 'fermentation_log', 'processing_sessions',
                                                    'hbts', 'filling_sessions')}
        print(data['batch_number'], data['type'] or 'unknown', 'HOLD' if data.get('hold') else counts,
              f'{len(data["warnings"])} warnings')
    if failed:
        print(f'{len(failed)} failed:', ', '.join(failed))


if __name__ == '__main__':
    main()
